from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from pypdf import PdfReader

from app.textutil import decode_text

Page = tuple[int | None, str]
SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx"}


class UnsupportedDocument(ValueError):
    pass


def load_document(filename: str, data: bytes) -> list[Page]:
    extension = Path(filename).suffix.lower()
    if extension in {".txt"}:
        return [(None, decode_text(data))]
    if extension in {".md", ".markdown"}:
        return [(None, _tidy_markdown(decode_text(data)))]
    if extension == ".pdf":
        return _load_pdf(data)
    if extension == ".docx":
        return _load_docx(data)
    raise UnsupportedDocument(f"Unsupported file type: {extension or 'unknown'}")


def _tidy_markdown(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


def _load_pdf(data: bytes) -> list[Page]:
    reader = PdfReader(BytesIO(data))
    pages: list[Page] = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append((number, text))
    return pages


def _load_docx(data: bytes) -> list[Page]:
    from docx import Document

    document = Document(BytesIO(data))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [(None, "\n\n".join(parts))]
