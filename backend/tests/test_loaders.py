from io import BytesIO

from docx import Document

from app.ingestion.loaders import UnsupportedDocument, load_document


def test_markdown_keeps_heading_text():
    pages = load_document("notes.md", b"# Refund policy\n\nBring the item back.")
    assert pages[0][0] is None
    assert "Refund policy" in pages[0][1]
    assert "Bring the item back." in pages[0][1]


def test_markdown_links_become_labels():
    pages = load_document("notes.md", b"Read the [packing slip](https://example.com/slip).")
    assert "packing slip" in pages[0][1]
    assert "https://example.com" not in pages[0][1]


def test_pdf_keeps_page_numbers():
    pages = load_document("policy.pdf", _pdf("Refunds are issued within 5 business days."))
    assert pages[0][0] == 1
    assert "5 business days" in pages[0][1]


def test_docx_reads_paragraphs_and_tables():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("The Halden lamp shade is frosted glass.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Bulb"
    table.rows[0].cells[1].text = "25000 hours"
    document.save(buffer)
    pages = load_document("faq.docx", buffer.getvalue())
    text = pages[0][1]
    assert "frosted glass" in text
    assert "Bulb" in text
    assert "25000 hours" in text


def test_unknown_extension_is_rejected():
    try:
        load_document("notes.csv", b"a,b")
    except UnsupportedDocument as exc:
        assert "csv" in str(exc)
    else:
        raise AssertionError("expected unsupported document")


def _pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)
